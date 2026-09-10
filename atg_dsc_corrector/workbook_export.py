"""Classeur de données avec grilles propres / Workbook with independent grids."""

from pathlib import Path
from pprint import pformat
import os
import math

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .i18n import tr as _t
from .models import same_path


def _column_header(name, units):
    text = str(name)
    unit = units.get(text)
    return f"{text} ({unit})" if unit and not text.endswith(f"({unit})") else text


def _fit_sheet(sheet, widths, heights, text_cells):
    """Use sizes collected during writing; only text can need extra wrapped rows."""
    for column, length in widths.items():
        dimension = sheet.column_dimensions[get_column_letter(column)]
        dimension.width = min(80, max(8, length + 3))
        dimension.bestFit = True
    for cell, span in text_cells:
        capacity = max(1, sum(sheet.column_dimensions[get_column_letter(col)].width
                              for col in range(cell.column, cell.column + span)) - 3)
        lines = sum(max(1, math.ceil(len(line) / capacity)) for line in cell.value.split("\n"))
        heights[cell.row] = max(heights[cell.row], 15 * lines + 4)
    for row in range(1, sheet.max_row + 1):
        sheet.row_dimensions[row].height = heights.get(row, 15)


def source_blocks(sources):
    blocks = []
    unique = []
    for source in sources:
        if source is None or any(
            same_path(source.source_path, other.source_path)
            and source.sheet_name == other.sheet_name
            and source.source_fingerprint == other.source_fingerprint
            for other in unique
        ):
            continue
        unique.append(source)
        from .exporters import _source_reference
        blocks.append((source.name, source.original_data(), {
            **_source_reference(source), "units": source.units,
            "column_mapping": source.mapping.as_dict(),
        }))
    return blocks


def write_workbook(destination, file_format, sheets, *, protected_paths=(), overwrite=False):
    # Older environments and oversized provenance comments keep the compatible writer.
    try:
        from .fast_workbook_export import write_fast_workbook
    except ModuleNotFoundError as exc:
        if exc.name != "xlsxwriter":
            raise
    else:
        if all(len(pformat(metadata, sort_dicts=False)) <= 32767
               for _title, blocks in sheets for _name, _frame, metadata in blocks):
            return write_fast_workbook(destination, file_format, sheets,
                                       protected_paths=protected_paths, overwrite=overwrite)
    return _write_openpyxl_workbook(destination, file_format, sheets,
                                     protected_paths=protected_paths, overwrite=overwrite)


def _write_openpyxl_workbook(destination, file_format, sheets, *, protected_paths=(), overwrite=False):
    from .exporters import ExportError, _temporary_path, _validate_destinations

    if file_format.lower().lstrip(".") != "xlsx":
        raise ExportError(_t("Le format d'export doit être XLSX."))
    target = Path(destination).with_suffix(".xlsx")
    _validate_destinations((target,), protected_paths, overwrite=overwrite)
    workbook = Workbook()
    alignment = Alignment(wrap_text=True, vertical="center")
    prototype = workbook.active.cell(1, 1)
    prototype.alignment = alignment
    body_style = prototype._style
    workbook.remove(workbook.active)
    try:
        for title, blocks in sheets:
            sheet = workbook.create_sheet(_t(title))
            sheet.freeze_panes = "A4"
            widths, heights, text_cells = {}, {}, []
            start = 1
            for name, frame, metadata in blocks:
                width = len(frame.columns)
                if not width:
                    continue
                if start + width - 1 > 16384 or len(frame) + 3 > 1048576:
                    raise ExportError(_t("Les données dépassent les dimensions maximales d'une feuille Excel."))
                for row, values in enumerate(
                    ([name], [str(metadata.get("source_path", ""))],
                     [_column_header(column, metadata.get("units", {})) for column in frame.columns]), start=1
                ):
                    for offset, value in enumerate(values):
                        cell = sheet.cell(row, start + offset, str(value))
                        cell.data_type = "s"
                        cell.alignment = alignment
                        heights[row] = 19
                        span = width if row < 3 else 1
                        text_cells.append((cell, span))
                        if span == 1:
                            lengths = [len(line) for line in cell.value.splitlines()]
                            widths[cell.column] = max(widths.get(cell.column, 0), max(lengths, default=0))
                        cell.font = Font(bold=True, color="FFFFFF" if row == 3 else "23394D")
                        if row == 3:
                            cell.fill = PatternFill("solid", fgColor="36566F")
                if width > 1:
                    for row in (1, 2):
                        sheet.merge_cells(start_row=row, start_column=start, end_row=row, end_column=start + width - 1)
                sheet.cell(1, start).comment = Comment(pformat(metadata, sort_dicts=False), "ThermalCurve")
                for row, values in enumerate(frame.itertuples(index=False, name=None), start=4):
                    for offset, value in enumerate(values):
                        if value is None or pd.isna(value):
                            continue
                        if isinstance(value, np.generic):
                            value = value.item()
                        if isinstance(value, float) and not np.isfinite(value):
                            continue
                        cell = sheet.cell(row, start + offset, value)
                        if cell.has_style:
                            cell.alignment = alignment  # Preserve automatic date/time number formats.
                        else:
                            cell._style = body_style  # Shared immutable style; no later cell style edits.
                        heights[row] = 19
                        if isinstance(value, str):
                            cell.data_type = "s"  # Source text must never become an Excel formula.
                            text_cells.append((cell, 1))
                            length = max(map(len, value.splitlines()), default=0)
                        else:
                            length = len(str(value))
                        widths[cell.column] = max(widths.get(cell.column, 0), length)
                sheet.column_dimensions[get_column_letter(start + width)].width = 3
                start += width + 1
            if start == 1:
                cell = sheet.cell(1, 1, _t("Aucune donnée."))
                cell.alignment = alignment
                widths[1], heights[1] = len(cell.value), 19
            _fit_sheet(sheet, widths, heights, text_cells)
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary = _temporary_path(target)
        try:
            workbook.save(temporary)
            os.replace(temporary, target)
        finally:
            temporary.unlink(missing_ok=True)
    finally:
        workbook.close()
    return target
